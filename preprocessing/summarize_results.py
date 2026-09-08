"""
Summarize Ablation Results
============================
Đọc TẤT CẢ file ablation_*.csv trong data/results/, gộp thành 1 file
Excel duy nhất (results_summary.xlsx), mỗi ablation 1 sheet, có định dạng
màu để dễ đọc — KHÔNG đụng, KHÔNG sửa bất kỳ file CSV gốc nào.

Usage:
    python summarize_results.py
    # → tạo data/results/results_summary.xlsx

An toàn 100% với CSV gốc: chỉ đọc (pd.read_csv), không bao giờ ghi đè.
"""

import os
import glob
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

RESULTS_DIR = os.path.join(os.path.dirname(__file__), "..", "data", "results")
OUT_PATH = os.path.join(RESULTS_DIR, "results_summary.xlsx")

HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
BEST_FILL = PatternFill(
    start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"
)  # xanh nhạt


def format_sheet(ws, df, highlight_col=None, minimize=True):
    """Định dạng 1 sheet: header màu, auto-width, highlight dòng RMSE tốt nhất."""
    for c_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=c_idx, value=col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for r_idx, row in enumerate(df.itertuples(index=False), start=2):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)

    # Auto width
    for c_idx, col_name in enumerate(df.columns, start=1):
        max_len = max(
            len(str(col_name)),
            df[col_name].map(lambda value: len(str(value))).max() if len(df) else 0,
        )
        ws.column_dimensions[get_column_letter(c_idx)].width = min(max_len + 3, 40)

    ws.freeze_panes = "A2"

    # Highlight dòng tốt nhất (RMSE nhỏ nhất hoặc lớn nhất tùy cột)
    if highlight_col and highlight_col in df.columns:
        idx_best = (
            df[highlight_col].idxmin() if minimize else df[highlight_col].idxmax()
        )
        for c_idx in range(1, len(df.columns) + 1):
            ws.cell(row=idx_best + 2, column=c_idx).fill = BEST_FILL


def main():
    csv_files = sorted(glob.glob(os.path.join(RESULTS_DIR, "ablation_*.csv")))
    if not csv_files:
        print(f"Không tìm thấy file ablation_*.csv nào trong {RESULTS_DIR}")
        return

    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)  # xóa sheet mặc định

    # Cột nào dùng để highlight "tốt nhất" cho từng loại ablation (tất cả đều minimize RMSE)
    RMSE_COL_CANDIDATES = ["rmse_mean", "RMSE_mean"]

    print(f"Tìm thấy {len(csv_files)} file kết quả:\n")
    for path in csv_files:
        fname = os.path.basename(path)
        sheet_name = fname.replace("ablation_", "").replace(".csv", "")[
            :31
        ]  # Excel giới hạn 31 ký tự
        df = pd.read_csv(path)
        print(f"  {fname:40s} -> sheet '{sheet_name}' ({len(df)} rows)")

        ws = wb.create_sheet(title=sheet_name)
        rmse_col = next((c for c in RMSE_COL_CANDIDATES if c in df.columns), None)
        format_sheet(ws, df, highlight_col=rmse_col, minimize=True)

    wb.save(OUT_PATH)
    print(f"\n✓ Đã lưu: {OUT_PATH}")
    print("  Mở file này để xem — hoàn toàn không ảnh hưởng tới các file CSV gốc.")
    print("  Dòng highlight xanh = RMSE tốt nhất trong mỗi bảng.")


if __name__ == "__main__":
    main()
